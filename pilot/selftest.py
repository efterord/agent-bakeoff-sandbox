"""Isolation self-tests and injected-defect negative controls (run on the runner
before any task; any failure aborts the pilot). All probes are harmless
negative controls: if one unexpectedly succeeded it would do nothing useful.
"""
from __future__ import annotations

import base64
import json
import os
import zipfile
from pathlib import Path

from pilot import sandbox
from pptxbench import checks_pptx, ooxml, scan
from pptxbench import design as D
from pptxbench import tasks as T

PROBE = ["/opt/pilot/venv/bin/python3", "/opt/pilot/harness/agent/probe.py"]
AGENT_MODEL = "claude-haiku-4-5-20251001"
EROFS = "refused:30"


def _probe(gate, image, work, name, mode, caps=None, with_gate=True):
    gdir = None
    if with_gate:
        gate.open(name, kind="agent", arm="probe", task="probe", caps=caps or {})
        gdir = Path(gate.host_dir(name))
    tdir = work / "selftest" / name
    sandbox.prepare(tdir, {"task.json": b"{}"})
    try:
        c = sandbox.run(f"pilot-{name}", image, PROBE + [mode], tdir / "in", tdir / "out", gdir, wall_s=300,
                        run_id=name)
        sandbox.release(tdir)
        probe = None
        p = tdir / "out" / "probe.json"
        if p.is_file() and not p.is_symlink():
            probe = json.loads(p.read_text())
        ex = scan.safe_extract(tdir / "out", work / "selftest" / (name + "-extracted"))
    finally:
        sandbox.unmount(tdir)
    g = gate.close(name) if with_gate else None
    return c, probe, ex, g


def _ledger(work):
    p = work / "gate" / "log" / "ledger.jsonl"
    return [json.loads(x) for x in p.read_text().splitlines() if x.strip()] if p.exists() else []


def run_all(gate, image, work: Path, out: Path, fixtures: dict, fx_render: dict, fx_keys: dict) -> dict:
    checks, failures = {}, []

    def expect(name, cond, detail=""):
        checks[name] = {"pass": bool(cond), "detail": str(detail)[:300]}
        if not cond:
            failures.append(f"{name}: {str(detail)[:200]}")

    # ---------------------------------------------------------- isolation
    c, p, _, g = _probe(gate, image, work, "selftest-iso", "isolation", caps={"max_requests": 3})
    expect("probe_ran", p is not None, (c.get("docker_run_err") or "") + c.get("log_tail", "")[-300:])
    p = p or {}
    eff = c.get("effective") or {}
    expect("hostconfig_network_none", eff.get("NetworkMode") == "none", eff.get("NetworkMode"))
    expect("hostconfig_readonly_rootfs", eff.get("ReadonlyRootfs") is True, eff.get("ReadonlyRootfs"))
    expect("hostconfig_cap_drop_all", eff.get("CapDrop") == ["ALL"] and not eff.get("CapAdd"), eff.get("CapDrop"))
    expect("hostconfig_not_privileged", eff.get("Privileged") is False, eff.get("Privileged"))
    expect("hostconfig_no_new_privileges", any("no-new-privileges" in s for s in eff.get("SecurityOpt") or []),
           eff.get("SecurityOpt"))
    expect("hostconfig_no_seccomp_override", not any("seccomp" in s for s in eff.get("SecurityOpt") or []),
           eff.get("SecurityOpt"))
    expect("container_removed_after_task", c.get("removed") is True, c.get("removed"))
    expect("non_root_uid", p.get("uid") == 10001 and p.get("gid") == 10001, (p.get("uid"), p.get("gid")))
    zero = "0000000000000000"
    expect("caps_dropped", p.get("cap_eff") == zero and p.get("cap_prm") == zero and p.get("cap_bnd") == zero,
           (p.get("cap_eff"), p.get("cap_bnd")))
    expect("no_new_privs_flag", p.get("no_new_privs") == "1", p.get("no_new_privs"))
    expect("seccomp_filter_active", p.get("seccomp") == "2", p.get("seccomp"))
    expect("read_only_root", all(v == EROFS for v in (p.get("rootfs_write") or {"x": "?"}).values()),
           p.get("rootfs_write"))
    expect("input_mount_read_only", p.get("in_write") == EROFS, p.get("in_write"))
    expect("output_mount_writable", p.get("out_write") == "WRITABLE", p.get("out_write"))
    expect("output_bounded_128m", 100 <= (p.get("out_size_mb") or 0) <= 130, p.get("out_size_mb"))
    expect("only_loopback_interface", p.get("net_ifaces") == ["lo"], p.get("net_ifaces"))
    expect("no_dns", str(p.get("dns", "")).startswith("failed"), p.get("dns"))
    expect("no_tcp_egress_testnet", str(p.get("tcp_testnet", "")).startswith("failed"), p.get("tcp_testnet"))
    expect("pids_limit_256", p.get("pids_max") == "256", p.get("pids_max"))
    expect("memory_limit_3g", p.get("memory_max") == str(3 * 1024 ** 3), p.get("memory_max"))
    expect("no_swap", p.get("memory_swap_max") in ("0", None), p.get("memory_swap_max"))
    expect("cpu_limit_1_5", p.get("cpu_max") == "150000 100000", p.get("cpu_max"))
    expect("fsize_limit", list(p.get("rlimit_fsize") or []) == [268435456, 268435456], p.get("rlimit_fsize"))
    expect("nofile_limit", list(p.get("rlimit_nofile") or []) == [1024, 1024], p.get("rlimit_nofile"))
    expect("no_core_dumps", list(p.get("rlimit_core") or []) == [0, 0], p.get("rlimit_core"))
    expect("no_secret_env", p.get("env_secret_like") == [] and p.get("env_value_keylike") == [],
           (p.get("env_secret_like"), p.get("env_value_keylike")))
    bad_env = [n for n in p.get("env_names") or [] if n.startswith(("GITHUB_", "ACTIONS_", "RUNNER_", "ANTHROPIC"))]
    expect("no_ci_env", not bad_env and p.get("env_names"), bad_env)
    expect("no_docker_socket", p.get("docker_sock") is False, p.get("docker_sock"))
    expect("gate_dir_only_socket", p.get("gate_dir") == ["run.sock"], p.get("gate_dir"))
    expect("gate_dir_read_only", p.get("gate_dir_write") == EROFS, p.get("gate_dir_write"))
    gg = p.get("gate") or {}
    expect("gate_reachable_via_ro_socket", gg.get("models", {}).get("status") == 200, gg.get("models"))
    expect("gate_model_overwritten", gg.get("plain", {}).get("status") == 200
           and gg.get("plain", {}).get("model") == AGENT_MODEL, gg.get("plain"))
    for k in ("server_tool", "url_image", "mcp", "keylike"):
        expect(f"gate_refuses_{k}", gg.get(k, {}).get("status") == 400, gg.get(k))
    for k in ("files_api", "batches"):
        expect(f"gate_forbids_{k}", gg.get(k, {}).get("status") == 403, gg.get(k))
    expect("gate_request_cap", [gg.get(f"cap_{i}", {}).get("status") for i in range(4)] == [200, 200, 400, 400]
           and gg.get("cap_2", {}).get("type") == "budget_exhausted", [gg.get(f"cap_{i}") for i in range(4)])
    led = [x for x in _ledger(work) if x.get("run") == "selftest-iso" and x.get("event") == "request"]
    first = led[0] if led else {}
    expect("gate_drops_beta_header", first.get("betas_dropped") == "probe-beta-2026-01-01", first.get("betas_dropped"))
    expect("gate_strips_sampling", "temperature" in (first.get("stripped") or []), first.get("stripped"))
    expect("gate_clamps_max_tokens", first.get("max_tokens") == 8192, first.get("max_tokens"))
    expect("gate_logs_requested_model", first.get("requested_model") == "claude-opus-5", first.get("requested_model"))
    expect("gate_run_summary", g and g.get("requests") == 3 and g.get("capped", 0) >= 2, g)

    # ---------------------------------------------- gate caps from the host
    gate.open("selftest-usd", caps={"max_usd": 0.000001})
    st, j = gate.messages("selftest-usd", {"model": "x", "max_tokens": 10, "messages": [{"role": "user", "content": "x"}]})
    gate.close("selftest-usd")
    expect("gate_usd_cap", st == 400 and (j.get("error") or {}).get("type") == "budget_exhausted", (st, j))
    gate.open("selftest-img", caps={"max_images": 2})
    imgs = [{"type": "image", "source": {"type": "base64", "media_type": "image/png",
                                         "data": base64.b64encode(bytes([i]) * 8).decode()}} for i in range(3)]
    st, j = gate.messages("selftest-img", {"model": "x", "max_tokens": 10, "messages": [{"role": "user", "content": imgs}]})
    gate.close("selftest-img")
    expect("gate_distinct_image_cap", st == 400 and (j.get("error") or {}).get("type") == "budget_exhausted", (st, j))
    gate.open("selftest-loosen", caps={"max_usd": 99.0, "max_requests": 9999})
    s2 = gate.close("selftest-loosen")
    expect("gate_caps_cannot_be_loosened", s2.get("caps", {}).get("max_usd") == 0.5
           and s2.get("caps", {}).get("max_requests") == 80, s2.get("caps"))

    # ------------------------------------------------ extraction controls
    c, _, ex, _ = _probe(gate, image, work, "selftest-plant", "plant", with_gate=False)
    refused = {x["name"]: x["reason"] for x in ex["refused"]}
    expect("extract_refuses_symlink", refused.get("deliverable.pptx") == "symlink", refused)
    expect("extract_refuses_hardlink", refused.get("canary.txt") == "hard link", refused)
    expect("extract_refuses_fifo", refused.get("trace/events.jsonl") == "fifo", refused)
    expect("extract_refuses_symlink_in_trace", refused.get("trace/env.jsonl") == "symlink", refused)
    expect("extract_refuses_unknown_name", "unexpected.bin" in refused, refused)
    expect("extract_copied_nothing", ex["copied"] == [], ex["copied"])
    c, _, ex, _ = _probe(gate, image, work, "selftest-bomb", "bomb", with_gate=False)
    bomb = work / "selftest" / "selftest-bomb-extracted" / "deliverable.pptx"
    fz = scan.zip_guard(bomb) if bomb.exists() else ["not extracted"]
    expect("zip_bomb_refused", any("ratio" in f or "uncompressed" in f for f in fz), fz)

    # ------------------------------------------- injected-defect fixtures
    nc = out / "negative-controls"
    nc.mkdir(parents=True, exist_ok=True)
    deck = next(p for k, p in fixtures.items() if k == "P3-harbor-en") if "P3-harbor-en" in fixtures \
        else next(iter(fixtures.values()))
    trunc = nc / "truncated.pptx"
    trunc.write_bytes(deck.read_bytes()[: len(deck.read_bytes()) // 2])
    expect("defect_truncated_zip", not ooxml.validate_package(trunc)["ok"] or scan.zip_guard(trunc))
    dangling = nc / "dangling.pptx"
    with zipfile.ZipFile(deck) as zi, zipfile.ZipFile(dangling, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            if item.filename != "ppt/slideLayouts/slideLayout6.xml":
                zo.writestr(item, zi.read(item.filename))
    v = ooxml.validate_package(dangling)
    expect("defect_dangling_relationship", not v["ok"] and any("dangling" in e for e in v["errors"]), v["errors"][:2])
    macro = nc / "macro.pptx"
    macro.write_bytes(deck.read_bytes())
    with zipfile.ZipFile(macro, "a") as z:
        z.writestr("ppt/vbaProject.bin", b"\0" * 16)
    expect("defect_macro_part", any("macro" in f for f in ooxml.security_scan(macro)))
    ext = nc / "external.pptx"
    with zipfile.ZipFile(deck) as zi, zipfile.ZipFile(ext, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                                    b'officeDocument/2006/relationships/oleObject" Target="http://203.0.113.7/x" '
                                    b'TargetMode="External"/></Relationships>')
            zo.writestr(item, data)
    expect("defect_external_relationship", any("external" in f for f in ooxml.security_scan(ext)),
           ooxml.security_scan(ext))
    trav = scan.make_traversal_zip(nc / "traversal.pptx")
    expect("defect_path_traversal_name", any("traversal" in f for f in scan.zip_guard(trav)))
    from pptx import Presentation
    tampered = nc / "notes-lost.pptx"
    prs = Presentation(str(deck))
    prs.slides[3].notes_slide.notes_text_frame.text = ""
    prs.slides[2].shapes[2].text_frame.paragraphs[0].runs[0].text = "tampered"
    prs.save(tampered)
    t6 = T.Task(tid="P6", klass="pptx-edit", fmt="pptx", kind="edit", fixture="deck", brief="",
                params={"scope": [2], "translate_slide": 2, "target": "ja", "notes_frozen": True}, design="harbor",
                lang="en")
    r = checks_pptx.evaluate(t6, deck, tampered, [], [])
    expect("defect_notes_lost_and_unrelated_text", any("notes" in x for x in r["critical"])
           and any("text changed" in x for x in r["critical"]) and not r["completed"], r["critical"])
    over = nc / "overflow.pptx"
    prs = Presentation(str(deck))
    prs.slides[1].shapes[1].text_frame.paragraphs[0].runs[0].text = D.CONTENT["en"]["mediocre_body"] * 3
    prs.save(over)
    s2 = [s for s in checks_pptx.inspect(over)["slides"] if s["idx"] == 2][0]
    expect("defect_overflow_estimate", any(x["overflow"] > 1.05 for x in s2["shapes"]))
    from pilot.run_pilot import render_sealed
    rr = render_sealed(image, {"overflow": over}, work, "nc-overflow")
    pages = rr.get("overflow", [])
    w, h = checks_pptx.inspect(over)["width"], checks_pptx.inspect(over)["height"]
    hits = checks_pptx._rendered_overflow(s2, pages[1] if len(pages) > 1 else None, w, h)
    expect("defect_overflow_confirmed_in_render", hits >= 1, f"pages={len(pages)} hits={hits}")
    offb = nc / "offbrand.pptx"
    prs = Presentation(str(deck))   # P5 fixture slide 6 is still Comic Sans / off-palette
    prs.save(offb)
    t5 = T.by_iid("P5-" + deck.stem)   # fixture decks are named <design>-<lang>.pptx
    r = checks_pptx.evaluate(t5, deck, offb, [], [])
    expect("defect_unredesigned_slide_off_brand", any("off-brand" in f for f in r["failures"]) and not r["completed"],
           r["failures"][:4])
    flagged = nc / "identifying.pptx"
    prs = Presentation(str(deck))
    prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = "built with Hermes pptx_create"
    prs.save(flagged)
    fl = scan.scrub_props(flagged, nc / "identifying-scrubbed.pptx")
    expect("blinding_flags_identifying_strings", any("Hermes" in f for f in fl), fl[:3])
    with zipfile.ZipFile(nc / "identifying-scrubbed.pptx") as z:
        core = z.read("docProps/core.xml")
    expect("blinding_scrubs_docprops", b"anonymous" in core and b"python-pptx" not in core)
    return {"checks": checks, "failures": failures, "passed": sum(v["pass"] for v in checks.values()),
            "total": len(checks)}

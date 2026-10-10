"""Deterministic PPTX checks: structure, design-system conformance, overflow,
alignment grid, charts/tables, speaker notes and out-of-scope regressions."""
from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path

from . import design as D
from . import imagediff, ooxml

TOL_EMU = int(0.06 * D.EMU_PER_INCH)
THEME_REFS = re.compile(r"^\+(mj|mn)-(lt|ea|cs)$")


def _char_em(ch: str) -> float:
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return 1.0
    if ch.isspace():
        return 0.28
    return 0.55 if ch.isalnum() else 0.4


def estimate_overflow(shape, default_pt: float = 18.0) -> float:
    """Estimated text height / frame height (>1 means the text does not fit).
    Uses per-character em widths (CJK = 1 em) and 1.2 line spacing."""
    from pptx.util import Pt
    if not shape.has_text_frame or not shape.width or not shape.height:
        return 0.0
    tf = shape.text_frame
    inset = Pt(7.2) * 2
    usable_w = max(1, shape.width - inset)
    total = 0.0
    for p in tf.paragraphs:
        size = None
        for r in p.runs:
            if r.font.size:
                size = r.font.size
                break
        size_emu = size if size else Pt(default_pt)
        text = "".join(r.text for r in p.runs) or " "
        width_emu = sum(_char_em(ch) for ch in text) * size_emu
        lines = max(1, math.ceil(width_emu / usable_w))
        total += lines * size_emu * 1.2
    return float(total + Pt(3.6) * 2) / float(shape.height)


def inspect(path: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(str(path))
    slides = []
    for idx, s in enumerate(prs.slides, start=1):
        shapes = []
        for sh in s.shapes:
            info = {"name": sh.name, "left": int(sh.left or 0), "top": int(sh.top or 0),
                    "width": int(sh.width or 0), "height": int(sh.height or 0), "text": "", "fonts": [],
                    "sizes": [], "colors": [], "overflow": 0.0, "chart": None, "table": None}
            if sh.has_text_frame:
                info["text"] = sh.text_frame.text
                info["overflow"] = estimate_overflow(sh)
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        rpr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                        if rpr is not None:
                            for tag in ("latin", "ea"):
                                el = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag)
                                if el is not None and el.get("typeface"):
                                    info["fonts"].append(el.get("typeface"))
                        if r.font.size:
                            info["sizes"].append(r.font.size.pt)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                info["chart"] = {
                    "categories": [str(c) for c in ch.plots[0].categories],
                    "series": {ser.name: list(ser.values) for ser in ch.plots[0].series},
                }
            if getattr(sh, "has_table", False) and sh.has_table:
                tbl = sh.table
                info["table"] = [[tbl.cell(r, c).text for c in range(len(tbl.columns))] for r in range(len(tbl.rows))]
                info["text"] = "\n".join(" ".join(row) for row in info["table"])
                # rendered height grows with rows; estimate it from row heights
                info["height"] = int(sum(int(r.height) for r in tbl.rows))
            xml = sh._element.xml
            info["colors"] = sorted(set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)))
            shapes.append(info)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        slides.append({"idx": idx, "id": s.slide_id, "layout": s.slide_layout.name, "shapes": shapes,
                       "notes": notes, "part": s.part.partname.lstrip("/"),
                       "text": "\n".join(x["text"] for x in shapes)})
    return {"slides": slides, "width": int(prs.slide_width), "height": int(prs.slide_height),
            "_emu": Emu}


def _fonts_ok(slide, design) -> float:
    allowed = D.allowed_fonts(design)
    fonts = [f for sh in slide["shapes"] for f in sh["fonts"]]
    if not fonts:
        return 1.0
    return sum(1 for f in fonts if f in allowed or THEME_REFS.match(f)) / len(fonts)


def _palette_ok(slide, design) -> float:
    pal = D.palette(design)
    cols = [c.upper() for sh in slide["shapes"] for c in sh["colors"]]
    if not cols:
        return 1.0
    return sum(1 for c in cols if c in pal) / len(cols)


def _grid_ok(slide, design, width_emu) -> float:
    lines = [g * D.EMU_PER_INCH for g in D.grid_lines_in(design, width_emu / D.EMU_PER_INCH)]
    shapes = [s for s in slide["shapes"] if s["width"] > 0]
    if not shapes:
        return 1.0
    good = 0
    for i, s in enumerate(shapes):
        others = [e for j, o in enumerate(shapes) if j != i for e in (o["left"], o["left"] + o["width"])]
        if any(abs(s["left"] - g) <= TOL_EMU for g in lines) or any(abs(s["left"] - e) <= TOL_EMU for e in others):
            good += 1
    return good / len(shapes)


def _bounds_ok(slide, w, h) -> bool:
    return all(s["left"] >= -TOL_EMU and s["top"] >= -TOL_EMU and s["left"] + s["width"] <= w + TOL_EMU
               and s["top"] + s["height"] <= h + TOL_EMU for s in slide["shapes"])


def _overlaps(slide) -> int:
    """Pairs of text-bearing shapes that overlap (cards behind text are fine)."""
    boxes = [s for s in slide["shapes"] if s["text"].strip()]
    n = 0
    for i, a in enumerate(boxes):
        for b in boxes[i + 1:]:
            ix = min(a["left"] + a["width"], b["left"] + b["width"]) - max(a["left"], b["left"])
            iy = min(a["top"] + a["height"], b["top"] + b["height"]) - max(a["top"], b["top"])
            if ix > TOL_EMU and iy > TOL_EMU:
                n += 1
    return n


def _rendered_overflow(slide, page_png, w, h) -> int:
    """Text frames whose estimate says overflow AND whose band below inks."""
    if page_png is None:
        return 0
    hits = 0
    for s in slide["shapes"]:
        if s["overflow"] > 1.05 and s["chart"] is None and s["table"] is None:
            box = (s["left"] / w, s["top"] / h, (s["left"] + s["width"]) / w, (s["top"] + s["height"]) / h)
            if imagediff.ink_below(page_png, box) > 0.01:
                hits += 1
    return hits


def _has_cjk(text: str) -> bool:
    return any(unicodedata.east_asian_width(ch) in ("W", "F") and not ch.isspace() for ch in text)


def evaluate(task, fixture: Path | None, out: Path, fx_pages: list[Path], out_pages: list[Path]) -> dict:
    design = D.DESIGNS[task.design]
    p = task.params
    res = {"checks": {}, "failures": [], "critical": [], "preservation": {}}
    chk, fail = res["checks"], res["failures"]
    o = inspect(out)
    w, h = o["width"], o["height"]
    slides = o["slides"]
    page_for = {s["idx"]: (out_pages[s["idx"] - 1] if len(out_pages) >= s["idx"] else None) for s in slides}

    # ---- design conformance on the slides the runtime authored
    if task.kind == "create":
        authored = slides
    else:
        fx = inspect(fixture)
        fx_ids = {s["id"] for s in fx["slides"]}
        scope_ids = {fx["slides"][i - 1]["id"] for i in p.get("scope", [])}
        authored = [s for s in slides if s["id"] in scope_ids or s["id"] not in fx_ids]
    if authored:
        chk["fonts"] = sum(_fonts_ok(s, design) for s in authored) / len(authored)
        chk["palette"] = sum(_palette_ok(s, design) for s in authored) / len(authored)
        chk["grid"] = sum(_grid_ok(s, design, w) for s in authored) / len(authored)
        est = [s for s in authored if any(x["overflow"] > 1.05 for x in s["shapes"]
                                          if x["chart"] is None and x["table"] is None)]
        rend = sum(_rendered_overflow(s, page_for.get(s["idx"]), w, h) for s in authored)
        chk["no_overflow"] = 1.0 - len(est) / len(authored)
        chk["rendered_overflow_frames"] = rend
        if est:
            fail.append(f"text overflow (estimate) on slides {[s['idx'] for s in est]}")
        if rend:
            fail.append(f"text overflow confirmed in render: {rend} frame(s)")
        chk["in_bounds"] = sum(1 for s in authored if _bounds_ok(s, w, h)) / len(authored)
        chk["no_text_overlap"] = sum(1 for s in authored if _overlaps(s) == 0) / len(authored)
    chk["rendered"] = 1.0 if out_pages else 0.0
    if not out_pages:
        fail.append("LibreOffice could not render the deck")

    # ---- task completion
    ok = True
    if task.kind == "create":
        n = len(slides)
        if not (p["min_slides"] <= n <= p["max_slides"]):
            ok = False
            fail.append(f"slide count {n} outside {p['min_slides']}-{p['max_slides']}")
        if p.get("notes_all") and not all(s["notes"].strip() for s in slides):
            ok = False
            fail.append("missing speaker notes")
        alltext = "\n".join(s["text"] for s in slides)
        for needle in p.get("must_contain", []):
            if needle not in alltext:
                ok = False
                fail.append(f"brief fact missing: {needle}")
        layouts = {s["layout"] for s in slides}
        fx = inspect(fixture) if fixture else None
        if fx and not layouts & {s["layout"] for s in fx["slides"]}:
            fail.append("does not use the template's layouts")
            ok = False
        res["completed"] = ok and chk.get("in_bounds", 0) == 1.0 and bool(out_pages) and \
            not chk.get("rendered_overflow_frames")
        return res

    fx_by_id = {s["id"]: s for s in fx["slides"]}
    out_by_id = {s["id"]: s for s in slides}
    scope_idx = set(p.get("scope", []))

    # out-of-scope slides: normalized XML identical, notes identical, render identical
    changed_parts = ooxml.diff_parts(fixture, out)
    regress, img_changed = [], []
    for fs in fx["slides"]:
        if fs["idx"] in scope_idx:
            continue
        osl = out_by_id.get(fs["id"])
        if osl is None:
            regress.append(f"slide {fs['idx']} deleted")
            res["critical"].append(f"unrelated slide {fs['idx']} deleted")
            continue
        # compare by slide id, not part name: python-pptx renames slide parts on save
        if ooxml.read_part(fixture, fs["part"]) != ooxml.read_part(out, osl["part"]):
            regress.append(f"slide {fs['idx']} XML changed")
            if fs["text"] != osl["text"]:
                res["critical"].append(f"unrelated slide {fs['idx']} text changed")
        if len(fx_pages) >= fs["idx"] and len(out_pages) >= osl["idx"]:
            m = imagediff.compare(fx_pages[fs["idx"] - 1], out_pages[osl["idx"] - 1])
            if m["changed"] > 0.005:
                img_changed.append({"slide": fs["idx"], **m})
    for fs in fx["slides"]:
        osl = out_by_id.get(fs["id"])
        if osl is not None and fs["notes"].strip() and osl["notes"] != fs["notes"]:
            regress.append(f"speaker notes of slide {fs['idx']} changed or lost")
            if p.get("notes_frozen") or fs["idx"] not in scope_idx or not osl["notes"].strip():
                res["critical"].append(f"speaker notes of slide {fs['idx']} lost/changed")
    design_parts = [n for n in changed_parts["changed"] + changed_parts["removed"]
                    if n.startswith(("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/"))]
    if design_parts:
        regress.append(f"master/layout/theme changed: {design_parts[:4]}")
    out_of_scope_slides = max(1, len(fx["slides"]) - len(scope_idx))
    res["preservation"] = {"regressions": regress, "image_changed": img_changed,
                           "design_parts_changed": design_parts}
    chk["preservation"] = max(0.0, 1.0 - (len(regress) + len(img_changed)) / out_of_scope_slides)
    chk["masters_preserved"] = 0.0 if design_parts else 1.0

    # in-scope requirements
    for sidx, needles in p.get("must_contain_slide", {}).items():
        osl = out_by_id.get(fx["slides"][sidx - 1]["id"])
        txt = osl["text"] if osl else ""
        for nd in needles:
            if nd not in txt:
                ok = False
                fail.append(f"slide {sidx} missing '{nd}'")
    for sidx, needles in p.get("must_not_contain_slide", {}).items():
        osl = out_by_id.get(fx["slides"][sidx - 1]["id"])
        for nd in needles:
            if osl and nd in osl["text"]:
                ok = False
                fail.append(f"slide {sidx} still contains old value '{nd}'")
    for sidx in p.get("keep_geometry", []):
        fs = fx["slides"][sidx - 1]
        osl = out_by_id.get(fs["id"])
        if osl is None:
            continue
        fg = sorted((s["left"], s["top"], s["width"]) for s in fs["shapes"])
        og = sorted((s["left"], s["top"], s["width"]) for s in osl["shapes"])
        same = len(fg) == len(og) and all(abs(a - b) <= TOL_EMU for x, y in zip(fg, og) for a, b in zip(x, y))
        chk[f"geometry_kept_s{sidx}"] = 1.0 if same else 0.0
        if not same:
            fail.append(f"slide {sidx} geometry changed (visual language not preserved)")
    if "insert_after" in p:
        pos = p["insert_after"]
        new = [s for s in slides if s["id"] not in fx_by_id]
        if len(slides) != len(fx["slides"]) + 1 or len(new) != 1:
            ok = False
            fail.append("expected exactly one new slide")
        else:
            ns = new[0]
            ref = fx["slides"][pos - 1]
            if ns["idx"] != pos + 1:
                ok = False
                fail.append(f"new slide at position {ns['idx']}, expected {pos + 1}")
            if ns["layout"] != ref["layout"]:
                fail.append("new slide does not reuse slide layout")
                ok = False
            if p["new_title"] not in ns["text"]:
                ok = False
                fail.append("new slide title missing")
            if sum(1 for b in p["new_bullets"] if b in ns["text"]) < len(p["new_bullets"]):
                ok = False
                fail.append("new slide bullets missing")
            if p.get("new_needs_notes") and not ns["notes"].strip():
                ok = False
                fail.append("new slide has no speaker notes")
    if "chart_slide" in p:
        osl = out_by_id.get(fx["slides"][p["chart_slide"] - 1]["id"])
        charts = [s["chart"] for s in (osl["shapes"] if osl else []) if s["chart"]]
        if not charts:
            ok = False
            fail.append("chart missing (native chart replaced or deleted)")
            chk["chart_integrity"] = 0.0
        else:
            ser = charts[0]["series"]
            good = all(abs((ser.get(k) or [0])[-1] - v) < 1e-6 for k, v in p["chart_q4"].items())
            cats = charts[0]["categories"] == D.QUARTERS
            prev_ok = all(list(ser.get(k, []))[:3] == D.SERIES[k][:3] for k in D.SERIES)
            chk["chart_integrity"] = (good + cats + prev_ok) / 3
            if not good:
                ok = False
                fail.append("chart Q4 values not updated")
    if "table_slide" in p:
        osl = out_by_id.get(fx["slides"][p["table_slide"] - 1]["id"])
        tables = [s for s in (osl["shapes"] if osl else []) if s["table"]]
        if not tables:
            ok = False
            fail.append("table missing")
            chk["table_integrity"] = 0.0
        else:
            t = tables[0]
            rows = t["table"]
            has = any(r and r[0] == p["new_row_first"] for r in rows)
            fits = t["top"] + t["height"] <= h - int(0.25 * D.EMU_PER_INCH)
            chk["table_integrity"] = (has + (len(rows) == p["table_rows"]) + fits) / 3
            if not has:
                ok = False
                fail.append("new table row missing")
            if not fits:
                fail.append("table runs into the footer band")
    if "improve_slide" in p:
        fs = fx["slides"][p["improve_slide"] - 1]
        osl = out_by_id.get(fs["id"])
        if osl is None:
            ok = False
        else:
            kw = sum(1 for k in p["keywords"] if k.lower() in osl["text"].lower())
            chk["message_kept"] = kw / len(p["keywords"])
            dens = max((len(s["text"]) for s in osl["shapes"]), default=0)
            limit = 140 if task.lang == "ja" else 260
            chk["density"] = 1.0 if dens <= limit else limit / dens
            if kw < len(p["keywords"]) - 1:
                ok = False
                fail.append("redesign dropped the slide's message")
            if _fonts_ok(osl, design) < 1.0 or _palette_ok(osl, design) < 0.9:
                ok = False
                fail.append("redesigned slide still off-brand")
    if "translate_slide" in p:
        fs = fx["slides"][p["translate_slide"] - 1]
        osl = out_by_id.get(fs["id"])
        body = "\n".join(s["text"] for s in (osl["shapes"] if osl else [])[1:])
        cjk = _has_cjk(body)
        if (p["target"] == "en" and cjk) or (p["target"] == "ja" and not cjk):
            ok = False
            fail.append("agenda not translated")
    if res["critical"]:
        ok = False
    res["completed"] = ok and bool(out_pages) and chk.get("in_bounds", 1.0) == 1.0 and \
        not chk.get("rendered_overflow_frames")
    return res

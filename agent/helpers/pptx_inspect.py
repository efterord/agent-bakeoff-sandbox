#!/usr/bin/env python3
"""Parity helper (pilot harness, MIT): structured inspection of a .pptx.

    pptx_inspect.py DECK --view shapes|layouts|theme

shapes : per slide (0-based index, slide id, layout) every shape with id, name,
         kind, placeholder idx, geometry in inches, text, fonts, sizes, colours.
layouts: master layout names (use them with pptx_from_template.py).
theme  : theme major/minor fonts and colour scheme.
"""
import argparse
import json
import re
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
IN = 914400.0


def shapes_view(prs):
    out = []
    for i, s in enumerate(prs.slides):
        items = []
        for sh in s.shapes:
            d = {"id": sh.shape_id, "name": sh.name, "kind": str(sh.shape_type).split(".")[-1].split(" ")[0],
                 "placeholder": sh.placeholder_format.idx if sh.is_placeholder else None,
                 "left": round((sh.left or 0) / IN, 2), "top": round((sh.top or 0) / IN, 2),
                 "width": round((sh.width or 0) / IN, 2), "height": round((sh.height or 0) / IN, 2)}
            if sh.has_text_frame:
                d["text"] = sh.text_frame.text[:500]
                fonts, sizes = set(), set()
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        rpr = r._r.find(A + "rPr")
                        if rpr is not None:
                            for tag in ("latin", "ea"):
                                el = rpr.find(A + tag)
                                if el is not None and el.get("typeface"):
                                    fonts.add(el.get("typeface"))
                        if r.font.size:
                            sizes.add(r.font.size.pt)
                d["fonts"], d["sizes"] = sorted(fonts), sorted(sizes)
            if getattr(sh, "has_table", False) and sh.has_table:
                d["table"] = [[sh.table.cell(r, c).text for c in range(len(sh.table.columns))]
                              for r in range(len(sh.table.rows))]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                d["chart"] = {"categories": [str(c) for c in sh.chart.plots[0].categories],
                              "series": {x.name: list(x.values) for x in sh.chart.plots[0].series}}
            d["colors"] = sorted(set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', sh._element.xml)))
            items.append(d)
        out.append({"index": i, "slide_id": s.slide_id, "layout": s.slide_layout.name, "shapes": items,
                    "notes": s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""})
    return {"slide_width_in": round(prs.slide_width / IN, 3), "slide_height_in": round(prs.slide_height / IN, 3),
            "slides": out}


def theme_view(prs):
    blob = prs.slide_master.part.part_related_by(RT.THEME).blob.decode("utf-8", "replace")
    fonts = re.findall(r'<a:(major|minor)Font>\s*<a:latin typeface="([^"]*)"', blob)
    ea = re.findall(r'<a:ea typeface="([^"]*)"', blob)
    colors = re.findall(r'<a:(dk1|lt1|dk2|lt2|accent\d|hlink|folHlink)>\s*<a:(?:srgbClr val|sysClr val="\w+" lastClr)="([0-9A-Fa-f]{6})"', blob)
    return {"fonts": {k: v for k, v in fonts}, "ea_fonts": sorted(set(x for x in ea if x)), "colors": dict(colors)}


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Inspect a .pptx (shapes, layouts, theme).")
    ap.add_argument("pptx")
    ap.add_argument("--view", choices=["shapes", "layouts", "theme"], required=True)
    a = ap.parse_args(argv)
    prs = Presentation(a.pptx)
    if a.view == "shapes":
        res = shapes_view(prs)
    elif a.view == "layouts":
        res = {"layouts": [{"index": i, "name": la.name, "placeholders": [p.placeholder_format.idx for p in la.placeholders]}
                           for i, la in enumerate(prs.slide_layouts)]}
    else:
        res = theme_view(prs)
    print(json.dumps(res, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())

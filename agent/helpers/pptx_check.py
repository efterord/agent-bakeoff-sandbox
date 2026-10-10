#!/usr/bin/env python3
"""Parity helper (pilot harness, MIT): deck checks the agents may run.

    pptx_check.py DECK --overflow   per text frame: estimated fill ratio (> 1.05 = overflow), out-of-bounds shapes
    pptx_check.py DECK --validate   package validity (zip, content types, relationships) + security findings

The overflow estimate is the same algorithm the harness scorer uses (per-character
em widths, CJK = 1 em, 1.2 line spacing). The scorer additionally confirms
overflow on the rendered page, which this helper does not do.
"""
import argparse
import json
import math
import posixpath
import sys
import unicodedata
import zipfile

IN = 914400.0
PT = 12700


def char_em(ch):
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return 1.0
    if ch.isspace():
        return 0.28
    return 0.55 if ch.isalnum() else 0.4


def estimate(shape, default_pt=18.0):
    if not shape.has_text_frame or not shape.width or not shape.height:
        return 0.0
    usable = max(1, shape.width - 2 * 7.2 * PT)
    total = 0.0
    for p in shape.text_frame.paragraphs:
        size = next((r.font.size for r in p.runs if r.font.size), None) or int(default_pt * PT)
        text = "".join(r.text for r in p.runs) or " "
        lines = max(1, math.ceil(sum(char_em(c) for c in text) * size / usable))
        total += lines * size * 1.2
    return float(total + 2 * 3.6 * PT) / float(shape.height)


def overflow(path):
    from pptx import Presentation
    prs = Presentation(path)
    w, h = prs.slide_width, prs.slide_height
    issues = []
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                continue
            r = estimate(sh)
            if r > 1.05:
                issues.append({"slide": i, "shape_id": sh.shape_id, "name": sh.name, "kind": "overflow",
                               "fill_ratio": round(r, 2)})
            l, t, sw, shh = sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0
            if l < -0.06 * IN or t < -0.06 * IN or l + sw > w + 0.06 * IN or t + shh > h + 0.06 * IN:
                issues.append({"slide": i, "shape_id": sh.shape_id, "name": sh.name, "kind": "out_of_bounds"})
    return {"ok": not issues, "issues": issues}


def validate(path):
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    errors, findings = [], []
    try:
        z = zipfile.ZipFile(path)
    except Exception as e:
        return {"ok": False, "errors": [f"not a zip: {e}"], "security": []}
    with z:
        names = set(z.namelist())
        if "[Content_Types].xml" not in names:
            errors.append("missing [Content_Types].xml")
        for n in sorted(names):
            low = n.lower()
            if low.endswith("vbaproject.bin"):
                findings.append(f"macro project {n}")
            if "/embeddings/" in low and low.endswith(".bin"):
                findings.append(f"OLE embedding {n}")
            if not (n.endswith(".xml") or n.endswith(".rels")):
                continue
            try:
                root = etree.fromstring(z.read(n), parser)
            except Exception as e:
                errors.append(f"malformed XML {n}: {e}")
                continue
            if n.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(n))
                for rel in root:
                    tgt = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        if "hyperlink" not in rel.get("Type", ""):
                            findings.append(f"external relationship -> {tgt}")
                        continue
                    full = posixpath.normpath(tgt.lstrip("/") if tgt.startswith("/") else posixpath.join(base, tgt))
                    if full not in names:
                        errors.append(f"dangling relationship {n} -> {tgt}")
    return {"ok": not errors and not findings, "errors": errors[:20], "security": findings[:20]}


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Check a .pptx for overflow or package problems.")
    ap.add_argument("pptx")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--overflow", action="store_true")
    g.add_argument("--validate", action="store_true")
    a = ap.parse_args(argv)
    res = overflow(a.pptx) if a.overflow else validate(a.pptx)
    print(json.dumps(res, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())

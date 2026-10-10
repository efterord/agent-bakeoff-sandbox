#!/usr/bin/env python3
"""Parity helper (pilot harness, MIT): shape-level edits on a .pptx.

    pptx_shape_ops.py DECK --ops OPS.json [--output OUT]

OPS.json is a list of {"op": ..., "slide": N (0-based), ...}. Shape selector
"shape": {"id": n} | {"name": s} | {"index": k} | {"text": exact} |
{"text_startswith": s} | {"placeholder": idx}. Units are inches; colours RRGGBB.

  set_text       shape, paragraphs[]      keeps each paragraph's first-run format
  set_font       shape, latin?, ea?, size?, color?, bold?
  set_fill       shape, color
  move_resize    shape, left?, top?, width?, height?
  add_textbox    left, top, width, height, paragraphs[], size, color?, bold?, latin?, ea?, bullet?
  add_shape      kind (rectangle|rounded_rectangle|oval), left, top, width, height, fill
  delete_shape   shape
  table_add_row  shape, values[]          clones the last row's formatting
  table_set_cell shape, row, col, text

All ops are applied in memory and saved once; any error exits non-zero with
the failing op index and nothing is written.
"""
import argparse
import copy
import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_IN = 914400
KINDS = {"rectangle": MSO_SHAPE.RECTANGLE, "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
         "oval": MSO_SHAPE.OVAL}


def emu(v):
    return Emu(int(round(float(v) * EMU_IN)))


def find(slide, sel):
    shapes = list(slide.shapes)
    if not isinstance(sel, dict) or len(sel) != 1:
        raise ValueError("shape selector must have exactly one key")
    (k, v), = sel.items()
    if k == "index":
        return shapes[int(v)]
    for sh in shapes:
        if k == "id" and sh.shape_id == int(v):
            return sh
        if k == "name" and sh.name == v:
            return sh
        if k == "placeholder" and sh.is_placeholder and sh.placeholder_format.idx == int(v):
            return sh
        if k in ("text", "text_startswith") and sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if (k == "text" and t == str(v).strip()) or (k == "text_startswith" and t.startswith(str(v))):
                return sh
    raise ValueError(f"no shape matches {sel}")


def set_ea(run, name):
    rpr = run._r.get_or_add_rPr()
    for old in rpr.findall(qn("a:ea")):
        rpr.remove(old)
    ea = rpr.makeelement(qn("a:ea"), {"typeface": name})
    latin = rpr.find(qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)
    else:
        rpr.append(ea)


def style(run, size=None, color=None, bold=None, latin=None, ea=None):
    if size is not None:
        run.font.size = Pt(float(size))
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bool(bold)
    if latin:
        run.font.name = latin
    if ea:
        set_ea(run, ea)


def op_set_text(slide, o):
    tf = find(slide, o["shape"]).text_frame
    lines = [str(x) for x in o["paragraphs"]]
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for i, p in enumerate(paras):
        if i < len(lines):
            if p.runs:
                p.runs[0].text = lines[i]
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
            else:
                p.add_run().text = lines[i]
        else:
            p._p.getparent().remove(p._p)


def op_set_font(slide, o):
    sh = find(slide, o["shape"])
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            style(r, o.get("size"), o.get("color"), o.get("bold"), o.get("latin"), o.get("ea"))


def op_set_fill(slide, o):
    sh = find(slide, o["shape"])
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(o["color"])


def op_move_resize(slide, o):
    sh = find(slide, o["shape"])
    for k in ("left", "top", "width", "height"):
        if o.get(k) is not None:
            setattr(sh, k, emu(o[k]))


def op_add_textbox(slide, o):
    box = slide.shapes.add_textbox(emu(o["left"]), emu(o["top"]), emu(o["width"]), emu(o["height"]))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, text in enumerate(o["paragraphs"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ("• " + str(text)) if o.get("bullet") else str(text)
        style(r, o.get("size", 18), o.get("color"), o.get("bold", False), o.get("latin"), o.get("ea"))


def op_add_shape(slide, o):
    sh = slide.shapes.add_shape(KINDS[o["kind"]], emu(o["left"]), emu(o["top"]), emu(o["width"]), emu(o["height"]))
    if o.get("fill"):
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(o["fill"])
    sh.line.fill.background()


def op_delete_shape(slide, o):
    el = find(slide, o["shape"])._element
    el.getparent().remove(el)


def op_table_add_row(slide, o):
    gf = find(slide, o["shape"])
    tbl = gf.table._tbl
    last = tbl.tr_lst[-1]
    new = copy.deepcopy(last)
    last.addnext(new)
    row = gf.table.rows[len(gf.table.rows) - 1]
    for ci, v in enumerate(o["values"]):
        para = row.cells[ci].text_frame.paragraphs[0]
        if para.runs:
            para.runs[0].text = str(v)
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            para.add_run().text = str(v)
    gf.height = gf.height + int(row.height)


def op_table_set_cell(slide, o):
    cell = find(slide, o["shape"]).table.cell(int(o["row"]), int(o["col"]))
    para = cell.text_frame.paragraphs[0]
    if para.runs:
        para.runs[0].text = str(o["text"])
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        para.add_run().text = str(o["text"])


OPS = {k[3:]: v for k, v in globals().items() if k.startswith("op_")}


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1], epilog=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--ops", required=True)
    ap.add_argument("--output")
    a = ap.parse_args(argv)
    with open(a.ops, encoding="utf-8") as fh:
        ops = json.load(fh)
    prs = Presentation(a.pptx)
    for i, o in enumerate(ops):
        try:
            fn = OPS[o["op"]]
            fn(prs.slides[int(o["slide"])], o)
        except Exception as e:  # report and write nothing
            print(json.dumps({"ok": False, "failed_op": i, "error": f"{type(e).__name__}: {e}"}))
            return 2
    out = a.output or a.pptx
    prs.save(out)
    print(json.dumps({"ok": True, "applied": len(ops), "output": out}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
